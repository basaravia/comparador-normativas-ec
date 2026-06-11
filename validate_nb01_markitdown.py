"""
Validation script for 01_markitdown_normativas.ipynb
Tests: markitdown import, PDF conversion, parse_normativa(), Excel/JSON export.
"""
import sys
import re
import json
import tempfile
from pathlib import Path

PROJECT_ROOT = Path(__file__).parent
PDF_PATH     = PROJECT_ROOT / "Normativa2026" / "PDL-DERECHOS-DIGITALES.pdf"

PASS = "[PASS]"
FAIL = "[FAIL]"

results = []


def check(label: str, condition: bool, detail: str = "") -> bool:
    status = PASS if condition else FAIL
    msg = f"{status}  {label}"
    if detail:
        msg += f"\n         detail: {detail}"
    print(msg)
    results.append((label, condition))
    return condition


# ── Test 1: imports ──────────────────────────────────────────────────────────
print("\n=== Test 1: imports ===")
try:
    from markitdown import MarkItDown
    import openpyxl
    check("markitdown importa sin error", True)
    check("openpyxl importa sin error", True)
except ImportError as e:
    check("imports críticos", False, str(e))
    sys.exit(1)


# ── Test 2: PDF existe ────────────────────────────────────────────────────────
print("\n=== Test 2: archivo PDF de prueba ===")
check(
    "PDF PDL-DERECHOS-DIGITALES.pdf existe",
    PDF_PATH.exists(),
    str(PDF_PATH),
)
if not PDF_PATH.exists():
    sys.exit(1)


# ── Test 3: conversión MarkItDown ────────────────────────────────────────────
print("\n=== Test 3: conversión MarkItDown (sin LLM) ===")
md = MarkItDown()
try:
    result = md.convert(str(PDF_PATH))
    texto = result.text_content
    check("md.convert() no lanza excepción", True)
    check(
        "texto extraído tiene > 1000 caracteres",
        len(texto) > 1000,
        f"{len(texto):,} chars extraídos",
    )
    check(
        "texto contiene 'Art' (indicativo de artículos)",
        "Art" in texto,
    )
except Exception as e:
    check("md.convert() no lanza excepción", False, str(e))
    texto = ""
    sys.exit(1)


# ── Test 4: parse_normativa ──────────────────────────────────────────────────
print("\n=== Test 4: parse_normativa() ===")

# Definición inline de parse_normativa (copiada del notebook)
_PAT_ART = re.compile(
    r'(?:^|\n)'
    r'(?:#{1,4}[ \t]+|[-*+][ \t]+|\d+\.[ \t]+)?'
    r'[ \t]{0,6}'
    r'Art(?:ículo|iculo|\.)[ \t]+'
    r'(\d+[\w]*)[ \t]*[.\-–—]?[ \t]*'
    r'([^\n]{0,250})',
    re.MULTILINE | re.IGNORECASE,
)
_PAT_JERARQUIA = re.compile(
    r'(?:^|\n)'
    r'(?:#{1,4}[ \t]+)?'
    r'[ \t]{0,4}'
    r'(TÍTULO|CAPÍTULO|SECCIÓN|Título|Capítulo|Sección)[ \t]+'
    r'([IVXLCDM\d]+|PRIMERO|SEGUNDO|TERCERO|CUARTO|QUINTO|SEXTO|'
    r'SÉPTIMO|OCTAVO|NOVENO|DÉCIMO|Único|ÚNICO)[ \t]*[.\-]?[ \t]*\n?'
    r'([^\n]{0,300})',
    re.MULTILINE,
)
_PAT_DISP = re.compile(
    r'(?:^|\n)[ \t]{0,4}'
    r'(DISPOSICIÓN(?:ES)?[ \t]+'
    r'(?:TRANSITORIA|GENERAL|FINAL|DEROGATORIA|REFORMATORIA|SUSTITUTIVA)S?)',
    re.MULTILINE | re.IGNORECASE,
)
_PAT_RESOL = re.compile(
    r'(?:^|\n)[ \t]{0,4}(CONSIDERANDO|RESUELVE|CERTIFICA|DISPONE)[ \t]*:',
    re.MULTILINE,
)
_PAT_ANEXO = re.compile(
    r'(?:^|\n)[ \t]{0,4}'
    r'(ANEXO[ \t]+(?:[IVXLCDM]+|\d+|Único|ÚNICO))[ \t]*[.\-]?[ \t]*\n?'
    r'([^\n]{0,400})',
    re.MULTILINE | re.IGNORECASE,
)
_PAT_FECHA = re.compile(
    r'\b(\d{1,2})\s+de\s+(enero|febrero|marzo|abril|mayo|junio|julio|agosto|'
    r'septiembre|octubre|noviembre|diciembre)\s+de\s+(\d{4})\b',
    re.IGNORECASE,
)


def _tipo_norma(texto: str) -> str:
    s = texto[:800].upper()
    for patron, tipo in [
        (r'RESOLUCIÓN', 'RESOLUCIÓN'),
        (r'PROYECTO DE LEY', 'PROYECTO DE LEY'),
        (r'LEY ORGÁNICA', 'LEY ORGÁNICA'),
        ('LEY', 'LEY'),
    ]:
        if re.search(patron, s):
            return tipo
    return 'NORMATIVA'


def _titulo_norma(texto: str) -> str:
    kws = ('LEY', 'RESOLUCIÓN', 'RESOLUCION', 'PROYECTO', 'REGLAMENTO', 'DECRETO')
    for linea in texto.split('\n')[:30]:
        l = linea.strip()
        if len(l) > 15 and any(k in l.upper() for k in kws):
            return l[:300]
    validas = [l.strip() for l in texto.split('\n') if l.strip()]
    return validas[0][:300] if validas else 'Sin título'


def _fecha_norma(texto: str) -> str:
    m = _PAT_FECHA.search(texto[:3000])
    if m:
        return f"{m.group(1)} de {m.group(2).lower()} de {m.group(3)}"
    return ""


def _seccion_de_articulo(pos: int, secciones: list) -> str:
    sec = None
    for s in secciones:
        if s['posicion'] <= pos:
            sec = s
        else:
            break
    if sec is None:
        return ""
    parts = [sec['tipo']]
    if sec['identificador']:
        parts.append(sec['identificador'])
    if sec['titulo']:
        parts.append(sec['titulo'][:80])
    return ' '.join(parts)


def parse_normativa(texto: str, archivo: str) -> dict:
    tipo   = _tipo_norma(texto)
    titulo = _titulo_norma(texto)
    fecha  = _fecha_norma(texto)
    am = [(m.group(1), (m.group(2) or '').strip(), m.start(), m.end())
          for m in _PAT_ART.finditer(texto)]
    secciones = []
    for m in _PAT_JERARQUIA.finditer(texto):
        tit = (m.group(3) or '').strip()
        tit = re.split(r'Art(?:ículo|\.)[ \t]+\d', tit)[0].strip()
        secciones.append({
            'tipo': m.group(1).upper(),
            'identificador': m.group(2).strip(),
            'titulo': tit[:250],
            'posicion': m.start(),
        })
    for m in _PAT_DISP.finditer(texto):
        secciones.append({'tipo': m.group(1).strip().upper(), 'identificador': '',
                          'titulo': '', 'posicion': m.start()})
    for m in _PAT_RESOL.finditer(texto):
        secciones.append({'tipo': m.group(1).upper(), 'identificador': '',
                          'titulo': '', 'posicion': m.start()})
    secciones.sort(key=lambda s: s['posicion'])
    articulos = []
    for i, (num, enc, start, end_m) in enumerate(am):
        next_pos  = am[i + 1][2] if i + 1 < len(am) else min(end_m + 4000, len(texto))
        contenido = re.sub(r'\n{3,}', '\n\n', texto[end_m:next_pos].strip())[:3500]
        seccion   = _seccion_de_articulo(start, secciones)
        articulos.append({
            'numero':     num,
            'encabezado': enc or None,
            'contenido':  contenido,
            'posicion':   start,
            'seccion':    seccion,
            'pagina':     None,
        })
    axm = list(_PAT_ANEXO.finditer(texto))
    anexos = []
    for i, m in enumerate(axm):
        c_start = m.end()
        c_end   = axm[i + 1].start() if i + 1 < len(axm) else min(c_start + 6000, len(texto))
        contenido = re.sub(r'\n{3,}', '\n\n', texto[c_start:c_end].strip())[:6000]
        anexos.append({
            'identificador': m.group(1).strip().upper(),
            'titulo':        (m.group(2) or '').strip(),
            'contenido':     contenido,
        })
    return {
        'archivo':     archivo,
        'titulo':      titulo,
        'tipo_norma':  tipo,
        'fecha':       fecha,
        'n_articulos': len(articulos),
        'n_secciones': len(secciones),
        'n_anexos':    len(anexos),
        'secciones':   secciones,
        'articulos':   articulos,
        'anexos':      anexos,
    }


try:
    doc = parse_normativa(texto, "PDL-DERECHOS-DIGITALES")
    check("parse_normativa() no lanza excepción", True)
    check(
        "n_articulos > 0",
        doc["n_articulos"] > 0,
        f"encontrados: {doc['n_articulos']}",
    )
    check(
        "n_articulos >= 30 (PDL-DERECHOS-DIGITALES tiene 39 segun notebook)",
        doc["n_articulos"] >= 30,
        f"encontrados: {doc['n_articulos']}",
    )
    check(
        "tipo_norma == 'PROYECTO DE LEY'",
        doc["tipo_norma"] == "PROYECTO DE LEY",
        f"tipo_norma={doc['tipo_norma']}",
    )
    check(
        "titulo contiene texto no vacío",
        len(doc["titulo"]) > 5,
        f"titulo='{doc['titulo'][:80]}'",
    )
    check(
        "articulos[0] tiene campo 'numero' no vacío",
        bool(doc["articulos"][0]["numero"]),
        f"numero='{doc['articulos'][0]['numero']}'",
    )
    check(
        "articulos[0] tiene campo 'contenido' no vacío",
        len(doc["articulos"][0].get("contenido", "")) > 0,
    )
    check(
        "n_secciones >= 0 (campo presente y válido)",
        isinstance(doc["n_secciones"], int) and doc["n_secciones"] >= 0,
        f"n_secciones={doc['n_secciones']}",
    )
except Exception as e:
    check("parse_normativa() no lanza excepción", False, str(e))
    sys.exit(1)


# ── Test 5: exportación JSON ─────────────────────────────────────────────────
print("\n=== Test 5: exportación JSON ===")
with tempfile.TemporaryDirectory() as tmpdir:
    json_path = Path(tmpdir) / "test_normativas.json"
    try:
        normativas = {"PDL-DERECHOS-DIGITALES": doc}
        with open(json_path, "w", encoding="utf-8") as f:
            json.dump(normativas, f, ensure_ascii=False, indent=2)
        check("json.dump() no lanza excepción", True)
        check("archivo JSON creado y no vacío", json_path.exists() and json_path.stat().st_size > 0)

        # Roundtrip: releer y verificar integridad
        with open(json_path, encoding="utf-8") as f:
            reload = json.load(f)
        check(
            "JSON roundtrip: n_articulos coincide",
            reload["PDL-DERECHOS-DIGITALES"]["n_articulos"] == doc["n_articulos"],
        )
    except Exception as e:
        check("exportación JSON", False, str(e))


# ── Test 6: exportación Excel ────────────────────────────────────────────────
print("\n=== Test 6: exportación Excel ===")
_ILLEGAL = re.compile(r'[\x00-\x08\x0b\x0c\x0e-\x1f]')


def _clean(v):
    return _ILLEGAL.sub('', v) if isinstance(v, str) else v


from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

with tempfile.TemporaryDirectory() as tmpdir:
    excel_path = Path(tmpdir) / "test_normativas.xlsx"
    try:
        wb = openpyxl.Workbook()
        wb.remove(wb.active)

        HDR_FONT  = Font(bold=True, color="FFFFFF", size=11)
        HDR_FILL  = PatternFill("solid", fgColor="2E75B6")
        HDR_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)
        CELL_TOP  = Alignment(vertical="top", wrap_text=True)
        THIN      = Side(style="thin", color="BFBFBF")
        BORDER    = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
        COLS      = [("NUMERO", 10), ("ARTICULO", 90), ("PAGINA", 10), ("SECCION", 45), ("FECHA", 22)]

        for nombre, d in normativas.items():
            ws = wb.create_sheet(title=nombre[:31])
            for col_i, (col_name, col_w) in enumerate(COLS, 1):
                c = ws.cell(row=1, column=col_i, value=col_name)
                c.font, c.fill, c.alignment, c.border = HDR_FONT, HDR_FILL, HDR_ALIGN, BORDER
                ws.column_dimensions[get_column_letter(col_i)].width = col_w
            ws.row_dimensions[1].height = 22
            ws.freeze_panes = "A2"

            fecha = d.get('fecha', '')
            for row_i, art in enumerate(d['articulos'], 2):
                texto_art = '\n'.join(filter(None, [art.get('encabezado'), art.get('contenido', '')])).strip()
                vals = [
                    art.get('numero', ''),
                    texto_art[:32767],
                    art.get('pagina') or 'N/D',
                    art.get('seccion', ''),
                    fecha,
                ]
                for col_i, val in enumerate(vals, 1):
                    c = ws.cell(row=row_i, column=col_i, value=_clean(val))
                    c.alignment = CELL_TOP
                    c.border    = BORDER
                ws.row_dimensions[row_i].height = 60

        wb.save(excel_path)
        check("openpyxl save() no lanza excepción", True)
        check("archivo Excel creado y no vacío", excel_path.exists() and excel_path.stat().st_size > 0)

        # Verificar que el Excel tiene la hoja y las filas esperadas
        wb2 = openpyxl.load_workbook(excel_path)
        ws2 = wb2["PDL-DERECHOS-DIGITALES"]
        # Fila 1 = headers, filas 2..N = artículos
        data_rows = ws2.max_row - 1
        check(
            "Excel contiene filas de datos igual a n_articulos",
            data_rows == doc["n_articulos"],
            f"filas={data_rows}, n_articulos={doc['n_articulos']}",
        )
        check(
            "Primera celda de header es 'NUMERO'",
            ws2.cell(1, 1).value == "NUMERO",
            f"valor={ws2.cell(1, 1).value}",
        )
    except Exception as e:
        check("exportación Excel", False, str(e))


# ── Resumen ───────────────────────────────────────────────────────────────────
print("\n" + "=" * 60)
passed = sum(1 for _, ok in results if ok)
failed = sum(1 for _, ok in results if not ok)
print(f"RESULTADO NB01: {passed} passed, {failed} failed  ({passed}/{len(results)})")
if failed:
    print("CHECKS FALLIDOS:")
    for label, ok in results:
        if not ok:
            print(f"  - {label}")
    sys.exit(1)
else:
    print("STATUS: PASS")
