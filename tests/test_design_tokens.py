"""Pruebas de `src/design_tokens.py` — fuente única de la paleta.

Tres cosas se verifican, alineadas con PLAN_MEJORAS_ANEXO.md §7 (ítem T) y §8:

1. Fuente única: ni `app/theme.py` ni `src/comparator.py` vuelven a escribir
   un hex de marca a mano (§2.2 defecto 5).
2. Fallback al placeholder cuando falta `assets/brand/brand.json`, para que
   un clon limpio del repositorio arranque sin nada del cliente.
3. Verificación automatizada de contraste WCAG 2.1: cada token cumple el
   mínimo declarado en `contraste_minimo`, y la regla cromo/dato (§8.2) —
   ningún estado del semáforo usa el primario ni el acento — se comprueba
   comparando los hex reales, no confiando en el dato de origen.

Todas las pruebas de contraste corren contra la fuente de tokens ACTIVA en
este entorno (la real si `assets/brand/brand.json` está presente, o el
placeholder si no) — no asumen cuál de las dos está cargada, salvo donde se
indica explícitamente.
"""
from __future__ import annotations

import json
import re
from pathlib import Path

import pytest

from src import design_tokens as tokens

HEX_LITERAL_RE = re.compile(r"#[0-9A-Fa-f]{6}\b")

THEME_PATH = tokens.REPO_ROOT / "app" / "theme.py"
COMPARATOR_PATH = tokens.REPO_ROOT / "src" / "comparator.py"


def _hex_literales(path: Path) -> set[str]:
    return {m.group(0).upper() for m in HEX_LITERAL_RE.finditer(path.read_text(encoding="utf-8"))}


# ── 1. Fuente única ──────────────────────────────────────────────────────────

def test_theme_no_tiene_hex_literales():
    encontrados = _hex_literales(THEME_PATH)
    assert not encontrados, (
        f"app/theme.py debería consumir src.design_tokens, no hex a mano: {sorted(encontrados)}"
    )


def test_comparator_no_tiene_hex_literales():
    encontrados = _hex_literales(COMPARATOR_PATH)
    assert not encontrados, (
        f"src/comparator.py debería consumir src.design_tokens, no hex a mano: {sorted(encontrados)}"
    )


def test_ningun_hex_duplicado_entre_theme_y_comparator():
    """Caso general (§2.2 defecto 5): aunque alguno de los dos volviera a
    tener hex a mano, no deberían coincidir entre sí — esa coincidencia
    accidental es exactamente el bug original."""
    assert not (_hex_literales(THEME_PATH) & _hex_literales(COMPARATOR_PATH))


# ── 2. Fallback al placeholder ───────────────────────────────────────────────

def test_fallback_al_placeholder_si_no_hay_brand_json(tmp_path):
    datos, ruta = tokens.load_tokens(
        brand_path=tmp_path / "brand.json",  # no existe
        placeholder_path=tokens.PLACEHOLDER_PATH,
    )
    assert ruta == tokens.PLACEHOLDER_PATH
    assert datos["meta"]["nombre"] == "Placeholder genérico"


def test_fallback_al_placeholder_si_brand_json_esta_corrupto(tmp_path):
    brand_roto = tmp_path / "brand.json"
    brand_roto.write_text("{ esto no es json", encoding="utf-8")
    _datos, ruta = tokens.load_tokens(brand_path=brand_roto, placeholder_path=tokens.PLACEHOLDER_PATH)
    assert ruta == tokens.PLACEHOLDER_PATH


def test_fallback_al_placeholder_si_brand_json_incompleto(tmp_path):
    brand_incompleto = tmp_path / "brand.json"
    brand_incompleto.write_text(json.dumps({"cromo": {}}), encoding="utf-8")
    _datos, ruta = tokens.load_tokens(brand_path=brand_incompleto, placeholder_path=tokens.PLACEHOLDER_PATH)
    assert ruta == tokens.PLACEHOLDER_PATH


def test_error_claro_si_ni_brand_ni_placeholder_existen(tmp_path):
    with pytest.raises(FileNotFoundError):
        tokens.load_tokens(brand_path=tmp_path / "a.json", placeholder_path=tmp_path / "b.json")


def test_brand_json_real_gana_si_esta_disponible():
    """No asume que `assets/brand/brand.json` exista (no se versiona), pero
    si este entorno lo tiene, debe preferirse sobre el placeholder."""
    if not tokens.BRAND_PATH.exists():
        pytest.skip("brand.json no está presente en este entorno (esperado en un clon limpio)")
    _datos, ruta = tokens.load_tokens()
    assert ruta == tokens.BRAND_PATH


def test_placeholders_de_logo_existen_con_los_nombres_exactos_de_8_4():
    """§8.4: brand.example.json referencia 5 logotipos placeholder; deben
    existir en disco con esos nombres exactos, independientemente de si el
    entorno además tiene la marca real (que no versiona sus propios logos)."""
    datos, ruta = tokens.load_tokens(
        brand_path=Path("/no/existe/brand.json"),
        placeholder_path=tokens.PLACEHOLDER_PATH,
    )
    assert ruta == tokens.PLACEHOLDER_PATH
    logos = {k: v for k, v in datos["logos"].items() if not k.startswith("_")}
    assert logos, "brand.example.json debería declarar al menos un logo"
    for nombre, ruta_relativa in logos.items():
        ruta_logo = tokens.REPO_ROOT / ruta_relativa
        assert ruta_logo.exists(), f"falta el placeholder de logo declarado en brand.example.json: {ruta_relativa}"
        assert ruta_logo.stat().st_size > 0, f"placeholder de logo vacío: {ruta_relativa}"


# ── 3. Contraste WCAG 2.1 ────────────────────────────────────────────────────

def test_ratio_contraste_casos_de_referencia():
    """Sanity check del algoritmo, independiente de cualquier marca."""
    assert tokens.ratio_contraste("#FFFFFF", "#000000") == pytest.approx(21.0, abs=0.01)
    assert tokens.ratio_contraste("#FFFFFF", "#FFFFFF") == pytest.approx(1.0, abs=0.001)
    assert tokens.ratio_contraste("#000000", "#000000") == pytest.approx(1.0, abs=0.001)
    # Simétrico: el orden de los argumentos no debe importar.
    assert tokens.ratio_contraste("#000000", "#FFFFFF") == tokens.ratio_contraste("#FFFFFF", "#000000")


@pytest.mark.parametrize("rol", ["primario", "primario_hover", "acento", "secundario"])
def test_ratio_cromo_coincide_con_lo_documentado_en_el_json(rol):
    """Los ratios que trae el propio brand.json/brand.example.json ("valores
    de referencia... documentados dentro de los propios JSON") deben poder
    reproducirse calculando desde el hex — si no coinciden, alguien cambió el
    color sin recalcular el ratio (o viceversa)."""
    nodo = tokens.CROMO[rol]
    calculado = tokens.ratio_contraste(nodo["hex"], "#FFFFFF")
    assert calculado == pytest.approx(nodo["ratio_sobre_blanco"], abs=0.05)


@pytest.mark.parametrize("rol", ["muted", "texto"])
def test_ratio_neutros_coincide_con_lo_documentado_en_el_json(rol):
    nodo = tokens.NEUTROS[rol]
    calculado = tokens.ratio_contraste(nodo["hex"], "#FFFFFF")
    assert calculado == pytest.approx(nodo["ratio_sobre_blanco"], abs=0.05)


@pytest.mark.parametrize("nivel", tokens.NIVEL_ORDEN)
def test_ratio_estado_coincide_con_lo_documentado_en_el_json(nivel):
    """El `ratio_texto` de cada estado es neutros.texto sobre el tinte —no
    sobre blanco—, tal como documenta `estado._doc` en el propio JSON."""
    nodo = tokens.ESTADO[nivel]
    calculado = tokens.ratio_contraste(tokens.TEXTO, nodo["tinte"])
    assert calculado == pytest.approx(nodo["ratio_texto"], abs=0.05)


def test_primario_cumple_texto_normal_para_titulares_y_botones():
    """`cromo.primario` se declara para "texto normal, titulares, fondo de
    botón con texto blanco" — debe superar el mínimo de texto normal, no
    solo el de componente UI."""
    ratio = tokens.ratio_contraste(tokens.PRIMARIO, "#FFFFFF")
    assert ratio >= tokens.CONTRASTE_MINIMO["texto_normal"]


def test_primario_hover_cumple_al_menos_componente_ui():
    """Solo se usa como relleno de área (§8.3): el mínimo exigible es el de
    componente UI, no el de texto normal."""
    ratio = tokens.ratio_contraste(tokens.PRIMARIO_HOVER, "#FFFFFF")
    assert ratio >= tokens.CONTRASTE_MINIMO["componente_ui"]


def test_secundario_cumple_texto_grande_y_componente_ui():
    """Declarado para "relleno, borde, icono, texto grande" — NO texto normal."""
    ratio = tokens.ratio_contraste(tokens.SECUNDARIO, "#FFFFFF")
    assert ratio >= tokens.CONTRASTE_MINIMO["texto_grande"]
    assert ratio >= tokens.CONTRASTE_MINIMO["componente_ui"]


def test_muted_y_texto_cumplen_texto_normal():
    assert tokens.ratio_contraste(tokens.MUTED, "#FFFFFF") >= tokens.CONTRASTE_MINIMO["texto_normal"]
    assert tokens.ratio_contraste(tokens.TEXTO, "#FFFFFF") >= tokens.CONTRASTE_MINIMO["texto_normal"]


@pytest.mark.parametrize("nivel", tokens.NIVEL_ORDEN)
def test_texto_sobre_tinte_de_estado_cumple_texto_normal(nivel):
    """El texto/etiqueta que se dibuja sobre cada tinte de estado (§8.2:
    "texto oscuro encima") debe seguir siendo legible. Si alguien mete un
    tinte demasiado oscuro o demasiado parecido al gris de texto, esta
    prueba debe fallar."""
    ratio = tokens.ratio_contraste(tokens.TEXTO, tokens.tinte(nivel))
    assert ratio >= tokens.CONTRASTE_MINIMO["texto_normal"]


# ── Colisión cromo/dato (§8.2) — la razón de ser de esta rama ───────────────

@pytest.mark.parametrize("nivel", tokens.NIVEL_ORDEN)
def test_estado_nunca_usa_cromo_primario_ni_acento(nivel):
    """La regla más importante del plan: ningún estado del semáforo puede
    coincidir con el verde institucional (`PRIMARIO`) ni con el acento lima
    (`ACENTO`) — si coincidiera, un "cumple" se confundiría con el cromo de
    la interfaz. Se compara el hex real, no se confía en que el JSON lo
    documente correctamente."""
    tinte_nivel = tokens.tinte(nivel).upper()
    assert tinte_nivel != tokens.PRIMARIO.upper(), f"estado.{nivel}.tinte coincide con cromo.primario"
    assert tinte_nivel != tokens.ACENTO.upper(), f"estado.{nivel}.tinte coincide con cromo.acento"


def test_nivel_colors_de_theme_nunca_usa_cromo():
    """Mismo chequeo, pero sobre la interfaz pública que consume
    `streamlit_app.py` (`NIVEL_COLORS`), para detectar una futura
    regresión introducida directamente en `app/theme.py`."""
    from app.theme import NIVEL_COLORS

    for nivel, colores in NIVEL_COLORS.items():
        for clave in ("fg", "bg"):
            valor = colores[clave].upper()
            assert valor != tokens.PRIMARIO.upper(), f"NIVEL_COLORS[{nivel!r}][{clave!r}] usa el primario de marca"
            assert valor != tokens.ACENTO.upper(), f"NIVEL_COLORS[{nivel!r}][{clave!r}] usa el acento de marca"


def test_acento_en_css_de_theme_nunca_es_color_de_texto():
    """§8.3: el acento no puede portar texto (2.20:1, falla AA). Verifica que
    toda línea del CSS inyectado por `inject_theme` que use la variable
    `--pb-accent` sea una propiedad de borde/decoración — nunca `color:`."""
    fuente = THEME_PATH.read_text(encoding="utf-8")
    usos = [linea for linea in fuente.splitlines() if "var(--pb-accent)" in linea]
    assert usos, "se esperaba al menos un uso de --pb-accent (border) para que este chequeo sea significativo"
    for linea in usos:
        propiedad = linea.strip().split(":", 1)[0]
        assert propiedad != "color", f"--pb-accent no puede ir en `color:` — línea: {linea!r}"


# ── Excel: el PatternFill sale de los mismos tokens que la UI ───────────────

def test_export_excel_usa_los_tokens_para_header_y_niveles(tmp_path):
    """Prueba de integración del DoD: "cambiar un token en un archivo cambia
    a la vez la UI y el Excel". Genera un Excel real con `export_excel` y
    verifica que los `PatternFill` resultantes son exactamente los hex de
    `design_tokens`, no valores hardcodeados en `src/comparator.py`."""
    import pandas as pd
    from openpyxl import load_workbook

    from src.comparator import DocumentComparator

    df = pd.DataFrame({
        "jerarquia": ["1.1", "1.2"],
        "nivel_cumplimiento": ["cumple", "omision"],
    })
    comparator = DocumentComparator(normativa_index=None, llm_grader=None)
    salida = comparator.export_excel(df, tmp_path / "reporte.xlsx")

    wb = load_workbook(salida)
    ws = wb["Comparación"]

    header_fill_real = ws[1][0].fill.start_color.rgb[-6:].upper()
    assert header_fill_real == tokens.hex_sin_almohadilla(tokens.PRIMARIO)

    fila_cumple = ws[2]  # primera fila de datos → nivel "cumple"
    fila_omision = ws[3]  # segunda fila de datos → nivel "omision"
    assert fila_cumple[0].fill.start_color.rgb[-6:].upper() == tokens.hex_sin_almohadilla(tokens.tinte("cumple"))
    assert fila_omision[0].fill.start_color.rgb[-6:].upper() == tokens.hex_sin_almohadilla(tokens.tinte("omision"))
